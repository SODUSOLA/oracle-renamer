"""
Read the staff sheet (.xlsx or .csv) into StaffRow objects.

Design notes
------------
* .xlsx is opened with openpyxl in `read_only=True` mode, which streams rows
  instead of loading the whole workbook DOM. Memory stays flat even for very
  large sheets.
* Column headers are matched case- and whitespace-insensitively, so
  "Oracle No", "ORACLE NO." and " oracle  no" all resolve to the same column.
* Every value passes through the same normalisers used by the indexer
  (see keys.py). That shared function is what makes the hash join correct.
"""

from __future__ import annotations

import csv
import logging
import re
from pathlib import Path
from typing import Dict, Iterable, Iterator, List, Optional, Sequence

from .keys import cell_to_text, normalise_key, normalise_lassra
from .models import StaffRow

log = logging.getLogger(__name__)


class ColumnNotFoundError(ValueError):
    """Raised when a required column is missing from the header row."""


def _header_key(text: object) -> str:
    """Reduce a header to letters+digits only: ' Oracle No. ' -> 'oracleno'."""
    return re.sub(r"[^a-z0-9]", "", str(text or "").lower())


def _resolve_columns(
    header: Sequence[object],
    oracle_col: str,
    lassra_col: str,
    name_col: Optional[str],
) -> Dict[str, Optional[int]]:
    """Find the index of each required column in the header row.

    Builds a dict {normalised header -> position} once (O(c) for c columns),
    then each lookup is O(1).
    """
    positions = {_header_key(h): i for i, h in enumerate(header) if _header_key(h)}

    def find(label: str, required: bool) -> Optional[int]:
        pos = positions.get(_header_key(label))
        if pos is None and required:
            raise ColumnNotFoundError(
                f"Column '{label}' not found. Available columns: "
                + ", ".join(str(h) for h in header if h is not None)
            )
        return pos

    return {
        "oracle": find(oracle_col, required=True),
        "lassra": find(lassra_col, required=True),
        "name": find(name_col, required=False) if name_col else None,
    }


def _iter_raw_rows(path: Path, sheet: Optional[str]) -> Iterator[Sequence[object]]:
    """Yield raw row tuples (header first) from an .xlsx or .csv file."""
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook  # imported lazily: CSV users don't need it

        # read_only=True streams rows; data_only=True returns formula RESULTS,
        # not formula text (a cell "=A2" should give its value).
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb[sheet] if sheet else wb.active
            yield from ws.iter_rows(values_only=True)
        finally:
            wb.close()  # read-only workbooks keep the file handle open until closed
    elif suffix == ".csv":
        # utf-8-sig strips the BOM Excel adds when saving "CSV UTF-8".
        with path.open(newline="", encoding="utf-8-sig") as fh:
            yield from csv.reader(fh)
    else:
        raise ValueError(f"Unsupported sheet type '{suffix}'. Use .xlsx or .csv")


def read_staff_rows(
    sheet_path: Path,
    oracle_col: str = "Oracle No",
    lassra_col: str = "LASSRA No",
    name_col: Optional[str] = "Name",
    sheet_name: Optional[str] = None,
    strip_leading_zeros: bool = False,
) -> List[StaffRow]:
    """Read and normalise every data row in the sheet.

    Rows are materialised into a list (O(n) memory) on purpose: the planner
    needs a full first pass to COUNT duplicate Oracle/LASSRA numbers before it
    can decide what to do with any single row (see ADR-005). A StaffRow is a
    handful of short strings, so even 500k rows is only tens of MB.
    """
    if not sheet_path.is_file():
        raise FileNotFoundError(f"Staff sheet not found: {sheet_path}")

    raw_rows = _iter_raw_rows(sheet_path, sheet_name)
    header = next(raw_rows, None)
    if header is None:
        raise ValueError(f"Staff sheet is empty: {sheet_path}")

    cols = _resolve_columns(header, oracle_col, lassra_col, name_col)
    rows: List[StaffRow] = []

    # enumerate from 2 because the header occupies spreadsheet row 1; this
    # lets users jump straight to the offending row in Excel.
    for row_number, raw in enumerate(raw_rows, start=2):
        if _is_blank(raw):
            continue  # trailing empty rows are common in Excel exports
        oracle_raw = cell_to_text(_get(raw, cols["oracle"]))
        lassra_raw = cell_to_text(_get(raw, cols["lassra"]))
        rows.append(
            StaffRow(
                row_number=row_number,
                oracle_raw=oracle_raw,
                lassra_raw=lassra_raw,
                name=cell_to_text(_get(raw, cols["name"])),
                oracle_key=normalise_key(oracle_raw, strip_leading_zeros),
                lassra_key=normalise_lassra(lassra_raw),
            )
        )

    log.info("Read %d staff row(s) from %s", len(rows), sheet_path.name)
    return rows


def _get(row: Sequence[object], index: Optional[int]) -> object:
    """Safe positional access: short rows (ragged CSV) return None."""
    if index is None or index >= len(row):
        return None
    return row[index]


def _is_blank(row: Iterable[object]) -> bool:
    return all(cell_to_text(cell) == "" for cell in row)
