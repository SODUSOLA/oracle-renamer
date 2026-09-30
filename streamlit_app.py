"""
Streamlit UI for oracle_renamer.

This file is a front end only: it wires uploaded files into a temp folder
and calls the exact same pipeline stages the CLI uses (indexer, reader,
planner, executor, report). No matching or copying logic lives here -- see
ADR-009 in the project doc, which anticipated a GUI replacing only the CLI
layer while the pipeline stays untouched.
"""

from __future__ import annotations

import io
import os
import shutil
import tempfile
import uuid
import zipfile
from pathlib import Path
from typing import List, Optional

import pandas as pd
import streamlit as st
from openpyxl import load_workbook

from oracle_renamer.executor import execute_plan
from oracle_renamer.indexer import DEFAULT_EXTENSIONS, build_image_index
from oracle_renamer.models import SUCCESS_STATUSES
from oracle_renamer.planner import build_plan
from oracle_renamer.reader import ColumnNotFoundError, read_staff_rows
from oracle_renamer.report import summarise, write_report

st.set_page_config(page_title="Oracle → LASSRA Photo Renamer", page_icon="", layout="wide")

RUNS_ROOT = Path(tempfile.gettempdir()) / "oracle_renamer_runs"


def _new_run_dir() -> Path:
    run_dir = RUNS_ROOT / uuid.uuid4().hex
    (run_dir / "images").mkdir(parents=True, exist_ok=True)
    (run_dir / "out").mkdir(parents=True, exist_ok=True)
    return run_dir


def _safe_extract_zip(zip_bytes: bytes, dest: Path) -> int:
    """Extract a zip into dest, refusing any entry that would escape it."""
    dest_resolved = dest.resolve()
    count = 0
    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
        for info in zf.infolist():
            if info.is_dir():
                continue
            target = (dest / info.filename).resolve()
            if dest_resolved != target and dest_resolved not in target.parents:
                continue  # zip-slip guard
            target.parent.mkdir(parents=True, exist_ok=True)
            with zf.open(info) as src, target.open("wb") as out:
                shutil.copyfileobj(src, out)
            count += 1
    return count


@st.cache_data(show_spinner=False)
def _sheet_names(file_bytes: bytes, filename: str) -> List[str]:
    """List worksheet names in an uploaded .xlsx, cached on file content."""
    if Path(filename).suffix.lower() not in {".xlsx", ".xlsm"}:
        return []
    fd, tmp_path = tempfile.mkstemp(suffix=Path(filename).suffix)
    os.close(fd)
    try:
        Path(tmp_path).write_bytes(file_bytes)
        wb = load_workbook(tmp_path, read_only=True)
        try:
            return wb.sheetnames
        finally:
            wb.close()
    finally:
        os.unlink(tmp_path)


def _zip_folder(folder: Path) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for path in sorted(folder.rglob("*")):
            if path.is_file() and path.suffix != ".part":
                zf.write(path, arcname=path.name)
    return buffer.getvalue()


st.title("Oracle → LASSRA Photo Renamer")
st.caption(
    "Upload Oracle-named staff photos and your staff sheet. Every row and "
    "every image is accounted for in a downloadable CSV report. Originals "
    "are never modified -- only copies get renamed."
)

if "result" not in st.session_state:
    st.session_state.result = None

with st.sidebar:
    st.header("1. Staff sheet")
    sheet_file = st.file_uploader("Excel or CSV", type=["xlsx", "xlsm", "csv"])

    sheet_name: Optional[str] = None
    if sheet_file is not None:
        names = _sheet_names(sheet_file.getvalue(), sheet_file.name)
        if len(names) > 1:
            picked = st.selectbox("Worksheet", ["(first sheet)"] + names)
            sheet_name = None if picked == "(first sheet)" else picked

    oracle_col = st.text_input("Oracle number column", "Oracle No")
    lassra_col = st.text_input("LASSRA number column", "LASSRA No")
    name_col = st.text_input("Name column (optional)", "Name")

    st.header("2. Photos")
    upload_mode = st.radio("Upload as", ["Individual files", "ZIP archive"], horizontal=True)
    image_files = None
    zip_file = None
    if upload_mode == "Individual files":
        image_files = st.file_uploader(
            "JPEG photos named by Oracle number", type=["jpg", "jpeg"],
            accept_multiple_files=True,
        )
    else:
        zip_file = st.file_uploader("ZIP of JPEG photos (subfolders OK)", type=["zip"])

    st.header("3. Options")
    with st.expander("Advanced", expanded=False):
        strip_leading_zeros = st.checkbox(
            "Treat 00123 and 123 as the same Oracle number", value=False,
        )
        recursive = st.checkbox("Scan subfolders (ZIP mode only)", value=True)
        dry_run = st.checkbox("Dry run (plan + report only, copy nothing)", value=False)
        overwrite = st.checkbox("Overwrite conflicting files at destination", value=False)
        verify_hash = st.checkbox("Verify with SHA-256 (slower, exact)", value=False)
        workers = st.slider("Parallel copy threads", 1, 16, 4)

    run_clicked = st.button("Run", type="primary", use_container_width=True)

if run_clicked:
    if sheet_file is None:
        st.error("Upload a staff sheet first.")
        st.stop()
    if upload_mode == "Individual files" and not image_files:
        st.error("Upload at least one photo first.")
        st.stop()
    if upload_mode == "ZIP archive" and zip_file is None:
        st.error("Upload a ZIP of photos first.")
        st.stop()

    run_dir = _new_run_dir()
    images_dir = run_dir / "images"
    out_dir = run_dir / "out"

    sheet_path = run_dir / f"staff{Path(sheet_file.name).suffix.lower()}"
    sheet_path.write_bytes(sheet_file.getvalue())

    if upload_mode == "Individual files":
        for f in image_files:
            (images_dir / Path(f.name).name).write_bytes(f.getvalue())
        effective_recursive = False
    else:
        _safe_extract_zip(zip_file.getvalue(), images_dir)
        effective_recursive = recursive

    try:
        with st.spinner("Indexing photos..."):
            index = build_image_index(
                images_dir, DEFAULT_EXTENSIONS, strip_leading_zeros, effective_recursive,
            )
        with st.spinner("Reading staff sheet..."):
            rows = read_staff_rows(
                sheet_path, oracle_col, lassra_col, name_col or None, sheet_name,
                strip_leading_zeros,
            )
        plan = build_plan(rows, index, out_dir)
        with st.spinner("Copying photos..."):
            execute_plan(plan, out_dir, dry_run, overwrite, verify_hash, workers)

        report_path = run_dir / "rename_report.csv"
        write_report(plan, report_path)
        summary = summarise(plan)
        all_ok = all(e.status in SUCCESS_STATUSES for e in plan.row_entries.values())
        report_bytes = report_path.read_bytes()

        st.session_state.result = {
            "summary": summary,
            "report_csv": report_bytes,
            "photos_zip": _zip_folder(out_dir) if not dry_run and plan.tasks else None,
            "report_df": pd.read_csv(io.BytesIO(report_bytes)),
            "all_ok": all_ok,
            "dry_run": dry_run,
        }
    except (FileNotFoundError, NotADirectoryError, ColumnNotFoundError, ValueError) as exc:
        st.error(str(exc))
        st.stop()
    finally:
        shutil.rmtree(run_dir, ignore_errors=True)

result = st.session_state.result
if result:
    st.divider()
    cols = st.columns(len(result["summary"]) or 1)
    for col, (status, count) in zip(cols, result["summary"].items()):
        col.metric(status, count)

    if result["all_ok"]:
        st.success("Every row is accounted for." + (" (dry run)" if result["dry_run"] else ""))
    else:
        st.warning("Some rows need attention -- see the report below.")

    st.subheader("Report")
    statuses = sorted(result["report_df"]["status"].unique())
    picked = st.multiselect("Filter by status", statuses, default=statuses)
    st.dataframe(
        result["report_df"][result["report_df"]["status"].isin(picked)],
        use_container_width=True, hide_index=True,
    )

    dl_cols = st.columns(2)
    dl_cols[0].download_button(
        "Download report (CSV)", result["report_csv"],
        file_name="rename_report.csv", mime="text/csv", use_container_width=True,
    )
    if result["photos_zip"]:
        dl_cols[1].download_button(
            "Download renamed photos (ZIP)", result["photos_zip"],
            file_name="renamed_photos.zip", mime="application/zip", use_container_width=True,
        )
    elif not result["dry_run"]:
        dl_cols[1].info("No photos were copied this run.")
else:
    st.info("Upload a staff sheet and photos, set your options, then click **Run**.")
